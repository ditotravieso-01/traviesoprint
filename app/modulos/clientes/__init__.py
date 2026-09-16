from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, send_file, current_app
from flask_login import login_required, current_user
from app.models import Client, User, Order, ArchivoAdjunto, Producto, Configuracion, EtiquetaFavorita
from app import db
from datetime import datetime, timedelta, date
import io
import os
import re
import uuid
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from werkzeug.utils import secure_filename
import json
from collections import Counter, defaultdict
from sqlalchemy import func

clientes_bp = Blueprint('clientes', __name__, url_prefix='/clientes', template_folder='templates')


# ==========================================
# CONSTANTES DE ARCHIVOS FAVORITOS
# ==========================================
ALLOWED_FAVORITO_TYPES = {
    'image/jpeg', 'image/jpg', 'image/png', 'image/gif', 'image/webp',
    'application/pdf'
}
DEFAULT_MAX_FAVORITO_MB = 50.0
PUNTOS_A_CM = 0.0352778  # 1 pt = 1/72 pulgada = 0.0352778 cm


# ==========================================
# DECORADORES
# ==========================================
def comercial_or_admin_required(func):
    from functools import wraps
    @wraps(func)
    def wrapper(*args, **kwargs):
        if current_user.role not in ['comercial', 'admin']:
            flash('No tienes permiso para realizar esta acción.', 'danger')
            return redirect(url_for('home.index'))
        return func(*args, **kwargs)
    return wrapper

def view_required(func):
    from functools import wraps
    @wraps(func)
    def wrapper(*args, **kwargs):
        if current_user.role not in ['comercial', 'admin', 'economico']:
            flash('No tienes permiso para ver clientes.', 'danger')
            return redirect(url_for('home.index'))
        return func(*args, **kwargs)
    return wrapper

def admin_required(func):
    from functools import wraps
    @wraps(func)
    def wrapper(*args, **kwargs):
        if current_user.role != 'admin':
            flash('Solo el administrador puede realizar esta acción.', 'danger')
            return redirect(url_for('clientes.listar_clientes'))
        return func(*args, **kwargs)
    return wrapper


# ==========================================
# CONFIGURACIÓN
# ==========================================
def obtener_config(clave, default=None):
    config = Configuracion.query.filter_by(clave=clave).first()
    if config:
        return config.valor
    return default

def guardar_config(clave, valor, descripcion=''):
    config = Configuracion.query.filter_by(clave=clave).first()
    if config:
        config.valor = str(valor)
        config.descripcion = descripcion
    else:
        config = Configuracion(clave=clave, valor=str(valor), descripcion=descripcion)
        db.session.add(config)
    db.session.commit()

def get_dashboard_config():
    default = {
        'graficos': {'nuevos_clientes': True, 'niveles': True, 'tipo_nuevos': 'bar', 'tipo_niveles': 'doughnut'},
        'widgets': {'nuevos_mes': True, 'retencion': True, 'top_clientes': True, 'recientes': True},
        'colores': {'grafico_nuevos': '#a8854f', 'grafico_niveles': ['#1a1a1a', '#d4a574', '#8e8e93']},
        'orden': ['kpis', 'graficos', 'widgets', 'recientes', 'niveles', 'acciones'],
        'estilo': {'altura_graficos': 80, 'mostrar_leyenda_niveles': True, 'mostrar_etiquetas': True}
    }
    try:
        config = json.loads(obtener_config('clientes_dashboard_config', '{}'))
        for key in default:
            if key not in config:
                config[key] = default[key]
            elif isinstance(default[key], dict):
                for subkey in default[key]:
                    if subkey not in config[key]:
                        config[key][subkey] = default[key][subkey]
        return config
    except:
        return default

def guardar_dashboard_config(config):
    guardar_config('clientes_dashboard_config', json.dumps(config), 'Configuración del dashboard de clientes')


# ==========================================
# HELPERS DE FAVORITOS
# ==========================================
def _get_max_favorito_mb():
    try:
        return float(obtener_config('clientes_favoritos_max_mb', str(DEFAULT_MAX_FAVORITO_MB)))
    except Exception:
        return DEFAULT_MAX_FAVORITO_MB


def _puede_subir_archivo_grande(user):
    """
    Devuelve True si el usuario puede saltarse el límite de tamaño.
    Por ahora: solo admin. Futuro: flag por usuario o permiso especial.
    """
    return bool(user and user.role == 'admin')


def _extraer_dimensiones_pdf(stream):
    """Intenta obtener el tamaño de la primera página en cm."""
    # 1. pypdf
    try:
        import pypdf
        stream.seek(0)
        reader = pypdf.PdfReader(stream)
        mb = reader.pages[0].mediabox
        return round(float(mb.width) * PUNTOS_A_CM, 2), round(float(mb.height) * PUNTOS_A_CM, 2)
    except Exception:
        pass
    # 2. PyPDF2
    try:
        import PyPDF2
        stream.seek(0)
        reader = PyPDF2.PdfReader(stream)
        mb = reader.pages[0].mediabox
        return round(float(mb.width) * PUNTOS_A_CM, 2), round(float(mb.height) * PUNTOS_A_CM, 2)
    except Exception:
        pass
    # 3. Regex sobre los primeros bytes
    try:
        stream.seek(0)
        data = stream.read(8192)
        m = re.search(rb'/MediaBox\s*\[\s*([\d.+-]+)\s+([\d.+-]+)\s+([\d.+-]+)\s+([\d.+-]+)\s*\]', data)
        if m:
            x1, y1, x2, y2 = [float(x) for x in m.groups()]
            w = abs(x2 - x1)
            h = abs(y2 - y1)
            return round(w * PUNTOS_A_CM, 2), round(h * PUNTOS_A_CM, 2)
    except Exception:
        pass
    return None, None


def _extraer_dimensiones_imagen(stream):
    """Lee dimensiones de la imagen. Usa DPI si está, si no asume 300 DPI."""
    try:
        from PIL import Image
        stream.seek(0)
        img = Image.open(stream)
        w_px, h_px = img.size
        dpi_info = img.info.get('dpi')
        dpi = 300.0
        if dpi_info and isinstance(dpi_info, (list, tuple)) and dpi_info[0] and dpi_info[0] > 0:
            dpi = float(dpi_info[0])
        return round(w_px / dpi * 2.54, 2), round(h_px / dpi * 2.54, 2)
    except Exception:
        return None, None


def _extraer_dimensiones_archivo(archivo):
    """Devuelve (ancho_cm, alto_cm) o (None, None) si no se pudo extraer."""
    if not archivo or not archivo.filename:
        return None, None
    mime = archivo.mimetype or ''
    try:
        archivo.seek(0)
        if mime == 'application/pdf':
            return _extraer_dimensiones_pdf(archivo)
        if mime.startswith('image/'):
            return _extraer_dimensiones_imagen(archivo)
    except Exception:
        pass
    finally:
        try:
            archivo.seek(0)
        except Exception:
            pass
    return None, None


def guardar_archivo_favorito(client_id, archivo, user):
    """
    Guarda el archivo en /static/uploads/favoritos/<client_id>/.
    Devuelve (ruta_relativa, nombre_original, mime, size_bytes) o (None, None, None, None).
    """
    if not archivo or not archivo.filename:
        return None, None, None, None

    mime = archivo.mimetype or ''
    if mime not in ALLOWED_FAVORITO_TYPES:
        raise ValueError('Tipo de archivo no permitido. Solo imágenes o PDF.')

    archivo.seek(0, os.SEEK_END)
    size = archivo.tell()
    archivo.seek(0)

    max_mb = _get_max_favorito_mb()
    max_bytes = int(max_mb * 1024 * 1024)
    if size > max_bytes and not _puede_subir_archivo_grande(user):
        raise ValueError(
            f'El archivo supera los {max_mb:.0f} MB permitidos. '
            f'Solo el administrador puede subir archivos de mayor tamaño.'
        )

    upload_folder = current_app.config.get('UPLOAD_FOLDER', 'app/static/uploads')
    favoritos_folder = os.path.join(upload_folder, 'favoritos', str(client_id))
    os.makedirs(favoritos_folder, exist_ok=True)

    nombre_original = secure_filename(archivo.filename)
    timestamp = datetime.now().strftime('%Y%m%d%H%M%S')
    uuid_part = str(uuid.uuid4())[:8]
    nombre_guardado = f"{timestamp}_{uuid_part}_{nombre_original}"
    ruta_relativa = os.path.join('favoritos', str(client_id), nombre_guardado)
    ruta_absoluta = os.path.join(upload_folder, ruta_relativa)
    archivo.save(ruta_absoluta)
    return ruta_relativa, nombre_original, mime, size


def borrar_archivo_favorito(fav):
    if not fav or not fav.archivo_ruta:
        return
    upload_folder = current_app.config.get('UPLOAD_FOLDER', 'app/static/uploads')
    ruta_absoluta = os.path.join(upload_folder, fav.archivo_ruta)
    try:
        if os.path.exists(ruta_absoluta):
            os.remove(ruta_absoluta)
    except Exception:
        pass


# ==========================================
# NIVELES / CONFIG / INSIGHTS
# ==========================================
def calcular_nivel_cliente(cliente, top_black_ids, top_golden_ids):
    if cliente.id in top_black_ids:
        return 'black'
    elif cliente.id in top_golden_ids:
        return 'golden'
    return 'standard'

def get_nivel_info(nivel, config_colores=None, config_nombres=None):
    colores_default = {'black': '#1a1a1a', 'golden': '#d4a574', 'standard': '#8e8e93'}
    nombres_default = {'black': 'Black', 'golden': 'Golden', 'standard': 'Standard'}
    if config_colores is None:
        try:
            config_colores = json.loads(obtener_config('clientes_colores_niveles', '{}'))
        except Exception:
            config_colores = {}
    if config_nombres is None:
        try:
            config_nombres = json.loads(obtener_config('clientes_nombres_niveles', '{}'))
        except Exception:
            config_nombres = {}
    color = config_colores.get(nivel, colores_default.get(nivel, '#8e8e93'))
    nombre = config_nombres.get(nivel, nombres_default.get(nivel, nivel.capitalize()))
    return {'color': color, 'nombre': nombre, 'label': nombre,
            'clase': f'nivel-{nivel}',
            'badge': '🔥' if nivel == 'black' else '⭐' if nivel == 'golden' else '👤'}

def get_columnas_visibles():
    try:
        columnas = json.loads(obtener_config('clientes_columnas_visibles', '[]'))
        if not columnas:
            columnas = ['referencia', 'nombre', 'telefono', 'email', 'pedidos', 'metros', 'nivel', 'acciones']
    except Exception:
        columnas = ['referencia', 'nombre', 'telefono', 'email', 'pedidos', 'metros', 'nivel', 'acciones']
    return columnas

def get_clasificacion_automatica():
    val = obtener_config('clientes_clasificacion_automatica', 'true')
    return val.lower() == 'true'

def get_clientes_nuevos_por_mes():
    ahora = datetime.now()
    meses, valores = [], []
    for i in range(11, -1, -1):
        mes = ahora.replace(day=1) - timedelta(days=30*i)
        meses.append(mes.strftime('%b %Y'))
        inicio = mes.replace(day=1).date()
        fin = datetime(mes.year + 1, 1, 1).date() if mes.month == 12 else datetime(mes.year, mes.month + 1, 1).date()
        valores.append(Client.query.filter(Client.created_at >= inicio, Client.created_at < fin).count())
    return meses, valores

def get_top_clientes_metros(limit=5):
    clientes = Client.query.all()
    for c in clientes:
        c.total_metros = calcular_metros_cliente(c)
    clientes.sort(key=lambda x: x.total_metros, reverse=True)
    top = clientes[:limit]
    return [c.nombre for c in top], [c.total_metros for c in top]

def get_tasa_retencion():
    total = Client.query.count()
    if total == 0:
        return 0
    repetidos = sum(1 for c in Client.query.all() if len(c.orders) > 1)
    return round((repetidos / total) * 100, 1)

def get_clientes_recientes(limit=5):
    return Client.query.order_by(Client.created_at.desc()).limit(limit).all()

def calcular_insights_cliente(cliente):
    orders = Order.query.filter_by(client_id=cliente.id).all()
    servicios_counter = Counter()
    materiales_counter = Counter()
    productos_counter = Counter()
    for o in orders:
        for s in o.get_servicios():
            servicios_counter[s] += 1
        for a in o.archivos:
            if a.nombre_visible:
                materiales_counter[a.nombre_visible] += 1
            if a.material:
                materiales_counter[a.material] += 1
            if a.producto_id:
                prod = Producto.query.get(a.producto_id)
                if prod:
                    productos_counter[prod.nombre] += 1
    return {
        'top_servicios': servicios_counter.most_common(5),
        'top_materiales': materiales_counter.most_common(5),
        'top_productos': productos_counter.most_common(5),
        'total_ordenes': len(orders),
        'frecuencia': calcular_frecuencia(orders)
    }

def calcular_frecuencia(orders):
    if len(orders) < 2:
        return 'eventual'
    fechas = [o.date.date() if isinstance(o.date, datetime) else o.date for o in orders if o.date]
    if len(fechas) < 2:
        return 'eventual'
    fechas.sort()
    diffs = [(fechas[i+1] - fechas[i]).days for i in range(len(fechas)-1)]
    avg = sum(diffs) / len(diffs)
    if avg <= 1: return 'diario'
    if avg <= 7: return 'semanal'
    if avg <= 14: return 'quincenal'
    if avg <= 31: return 'mensual'
    if avg <= 92: return 'trimestral'
    return 'eventual'

def calcular_metros_cliente(cliente):
    total = 0.0
    for order in cliente.orders:
        for archivo in order.archivos:
            if archivo.parametros_etiqueta:
                try:
                    params = json.loads(archivo.parametros_etiqueta)
                    if params.get('consumo'):
                        total += params.get('consumo', 0)
                        continue
                except Exception:
                    pass
            if archivo.unidad in ['m', 'm2', 'm²'] and archivo.cantidad:
                total += archivo.cantidad
    return round(total, 2)


# ==========================================
# LISTAR CLIENTES
# ==========================================
@clientes_bp.route('/')
@login_required
@view_required
def listar_clientes():
    search = request.args.get('search', '').strip()
    sort = request.args.get('sort', 'metros')
    order = request.args.get('order', 'desc')
    page = request.args.get('page', 1, type=int)
    per_page = 20

    query = Client.query
    if search:
        query = query.filter(db.or_(
            Client.nombre.ilike(f'%{search}%'),
            Client.referencia.ilike(f'%{search}%'),
            Client.email.ilike(f'%{search}%'),
            Client.telefono.ilike(f'%{search}%')
        ))

    clientes_todos = query.all()
    for c in clientes_todos:
        c.total_metros = calcular_metros_cliente(c)
        c.total_pedidos = len(c.orders) if c.orders else 0
        ins = calcular_insights_cliente(c)
        c.top_materiales = ins['top_materiales'][:2]
        c.top_servicios = ins['top_servicios'][:2]
        c.frecuencia = ins['frecuencia']

    clasificacion_automatica = get_clasificacion_automatica()
    try:
        config_colores = json.loads(obtener_config('clientes_colores_niveles', '{}'))
    except Exception:
        config_colores = {}
    try:
        config_nombres = json.loads(obtener_config('clientes_nombres_niveles', '{}'))
    except Exception:
        config_nombres = {}

    if clasificacion_automatica and len(clientes_todos) >= 5:
        try:
            top_black_pct = float(obtener_config('clientes_top_black_pct', '5'))
            top_golden_pct = float(obtener_config('clientes_top_golden_pct', '20'))
        except Exception:
            top_black_pct, top_golden_pct = 5, 20
        s = sorted(clientes_todos, key=lambda c: c.total_metros, reverse=True)
        n = len(s)
        top_black_ids = set(c.id for c in s[:max(1, int(n * top_black_pct / 100))])
        top_golden_ids = set(c.id for c in s[:max(1, int(n * top_golden_pct / 100))])
    else:
        top_black_ids, top_golden_ids = set(), set()

    nivel_map = {}
    black_count = golden_count = standard_count = 0
    for c in clientes_todos:
        nivel = calcular_nivel_cliente(c, top_black_ids, top_golden_ids) if clasificacion_automatica else 'standard'
        nivel_map[c.id] = nivel
        if nivel == 'black': black_count += 1
        elif nivel == 'golden': golden_count += 1
        else: standard_count += 1
        c.nivel = nivel
        c.nivel_info = get_nivel_info(nivel, config_colores, config_nombres)

    key_funcs = {
        'referencia': lambda x: x.referencia,
        'nombre': lambda x: x.nombre,
        'telefono': lambda x: x.telefono or '',
        'email': lambda x: x.email or '',
        'pedidos': lambda x: x.total_pedidos,
        'metros': lambda x: x.total_metros,
        'nivel': lambda x: {'black': 3, 'golden': 2, 'standard': 1}.get(x.nivel, 0),
    }
    key_func = key_funcs.get(sort, lambda x: x.total_metros)
    clientes_todos.sort(key=key_func, reverse=(order == 'desc'))

    total = len(clientes_todos)
    total_pages = (total + per_page - 1) // per_page
    clients = clientes_todos[(page - 1) * per_page: page * per_page]

    total_ordenes = sum(c.total_pedidos for c in clientes_todos)
    fecha_limite = datetime.now().date() - timedelta(days=30)
    clientes_activos = 0
    for c in clientes_todos:
        for o in c.orders:
            if o.date and o.date >= fecha_limite:
                clientes_activos += 1
                break

    top_cliente = max(clientes_todos, key=lambda x: x.total_metros) if clientes_todos else None

    meses_nuevos, valores_nuevos = get_clientes_nuevos_por_mes()
    top_nombres, top_metros = get_top_clientes_metros(5)
    tasa_retencion = get_tasa_retencion()
    clientes_recientes = get_clientes_recientes(5)

    nivel_labels, nivel_data, nivel_colors = [], [], []
    for nivel in ['black', 'golden', 'standard']:
        nivel_labels.append(config_nombres.get(nivel, nivel.capitalize()))
        nivel_data.append(sum(1 for c in clientes_todos if nivel_map.get(c.id) == nivel))
        nivel_colors.append(config_colores.get(nivel, '#8e8e93'))

    return render_template('clientes.html',
        clients=clients, search=search, sort=sort, order=order,
        page=page, per_page=per_page, total=total, total_pages=total_pages,
        total_clientes=total, total_ordenes=total_ordenes,
        clientes_activos=clientes_activos,
        top_cliente_nombre=top_cliente.nombre if top_cliente else '—',
        top_cliente_pedidos=top_cliente.total_pedidos if top_cliente else 0,
        top_cliente_metros=top_cliente.total_metros if top_cliente else 0,
        black_count=black_count, golden_count=golden_count, standard_count=standard_count,
        columnas_visibles=get_columnas_visibles(),
        clasificacion_automatica=clasificacion_automatica,
        config_colores=config_colores, config_nombres=config_nombres,
        meses_nuevos=meses_nuevos, valores_nuevos=valores_nuevos,
        top_nombres=top_nombres, top_metros=top_metros,
        top_clientes_data=list(zip(top_nombres, top_metros)),
        tasa_retencion=tasa_retencion, clientes_recientes=clientes_recientes,
        nivel_labels=nivel_labels, nivel_data=nivel_data, nivel_colors=nivel_colors,
        dashboard_config=get_dashboard_config())


# ==========================================
# CONFIGURACIÓN
# ==========================================
@clientes_bp.route('/configuracion', methods=['GET'])
@login_required
@admin_required
def get_configuracion():
    try:
        colores = json.loads(obtener_config('clientes_colores_niveles', '{}'))
    except Exception:
        colores = {}
    try:
        nombres = json.loads(obtener_config('clientes_nombres_niveles', '{}'))
    except Exception:
        nombres = {}
    try:
        columnas = json.loads(obtener_config('clientes_columnas_visibles', '[]'))
    except Exception:
        columnas = []
    return jsonify({
        'top_black_pct': float(obtener_config('clientes_top_black_pct', '5')),
        'top_golden_pct': float(obtener_config('clientes_top_golden_pct', '20')),
        'columnas_visibles': columnas,
        'colores': colores,
        'nombres': nombres,
        'clasificacion_automatica': obtener_config('clientes_clasificacion_automatica', 'true') == 'true',
        'dashboard_config': get_dashboard_config(),
        'favoritos_max_mb': _get_max_favorito_mb(),
    })

@clientes_bp.route('/configuracion', methods=['POST'])
@login_required
@admin_required
def set_configuracion():
    data = request.get_json() or {}
    if 'top_black_pct' in data and 'top_golden_pct' in data:
        try:
            tb, tg = float(data['top_black_pct']), float(data['top_golden_pct'])
            if tb <= 0 or tg <= 0 or tb >= tg:
                return jsonify({'error': 'Los valores deben ser positivos y Black < Golden'}), 400
            guardar_config('clientes_top_black_pct', str(tb))
            guardar_config('clientes_top_golden_pct', str(tg))
        except Exception:
            return jsonify({'error': 'Valores inválidos'}), 400
    if 'columnas_visibles' in data:
        guardar_config('clientes_columnas_visibles', json.dumps(data['columnas_visibles']))
    if 'colores' in data:
        guardar_config('clientes_colores_niveles', json.dumps(data['colores']))
    if 'nombres' in data:
        guardar_config('clientes_nombres_niveles', json.dumps(data['nombres']))
    if 'clasificacion_automatica' in data:
        guardar_config('clientes_clasificacion_automatica', 'true' if data['clasificacion_automatica'] else 'false')
    if 'dashboard_config' in data:
        guardar_dashboard_config(data['dashboard_config'])
    if 'favoritos_max_mb' in data:
        try:
            mb = float(data['favoritos_max_mb'])
            if mb > 0:
                guardar_config('clientes_favoritos_max_mb', str(mb))
        except Exception:
            pass
    return jsonify({'success': True})


# ==========================================
# ADMIN PANEL
# ==========================================
@clientes_bp.route('/admin')
@login_required
@admin_required
def admin_panel():
    total_clientes = Client.query.count()
    total_ordenes = Order.query.count()
    clientes_todos = Client.query.all()
    for c in clientes_todos:
        c.total_metros = calcular_metros_cliente(c)
    try:
        colores = json.loads(obtener_config('clientes_colores_niveles', '{}'))
    except Exception:
        colores = {}
    try:
        nombres = json.loads(obtener_config('clientes_nombres_niveles', '{}'))
    except Exception:
        nombres = {}
    columnas = get_columnas_visibles()
    clasificacion_auto = get_clasificacion_automatica()

    black_count = golden_count = 0
    if clasificacion_auto and len(clientes_todos) >= 5:
        try:
            tb = float(obtener_config('clientes_top_black_pct', '5'))
            tg = float(obtener_config('clientes_top_golden_pct', '20'))
        except Exception:
            tb, tg = 5, 20
        s = sorted(clientes_todos, key=lambda c: c.total_metros, reverse=True)
        n = len(s)
        black_ids = set(c.id for c in s[:max(1, int(n * tb / 100))])
        golden_ids = set(c.id for c in s[:max(1, int(n * tg / 100))])
        black_count = sum(1 for c in clientes_todos if c.id in black_ids)
        golden_count = sum(1 for c in clientes_todos if c.id in golden_ids)

    return render_template('admin_clientes.html',
        top_black_pct=obtener_config('clientes_top_black_pct', '5'),
        top_golden_pct=obtener_config('clientes_top_golden_pct', '20'),
        total_clientes=total_clientes, total_ordenes=total_ordenes,
        black_count=black_count, golden_count=golden_count,
        standard_count=total_clientes - black_count - golden_count,
        colores=colores, nombres=nombres,
        columnas_visibles=columnas,
        columnas_disponibles=[
            {'id': 'referencia', 'label': 'Referencia'},
            {'id': 'nombre', 'label': 'Nombre'},
            {'id': 'telefono', 'label': 'Teléfono'},
            {'id': 'email', 'label': 'Email'},
            {'id': 'pedidos', 'label': 'Pedidos'},
            {'id': 'metros', 'label': 'Metros'},
            {'id': 'nivel', 'label': 'Nivel'},
            {'id': 'acciones', 'label': 'Acciones'},
        ],
        clasificacion_automatica=clasificacion_auto,
        dashboard_config=get_dashboard_config())


# ==========================================
# CREAR CLIENTE
# ==========================================
@clientes_bp.route('/crear', methods=['GET', 'POST'])
@login_required
@comercial_or_admin_required
def crear_cliente():
    if request.method == 'POST':
        nombre = request.form.get('nombre', '').strip()
        referencia = request.form.get('referencia', '').strip()
        if not nombre:
            flash('El nombre es obligatorio.', 'danger')
            return render_template('form_cliente.html')
        if not referencia:
            referencia = f"CLI-{datetime.now().strftime('%Y%m%d%H%M%S')}"
        elif Client.query.filter_by(referencia=referencia).first():
            flash(f'Ya existe un cliente con la referencia "{referencia}".', 'danger')
            return render_template('form_cliente.html')

        cliente = Client(
            referencia=referencia,
            nombre=nombre,
            telefono=request.form.get('telefono', '').strip(),
            email=request.form.get('email', '').strip(),
            direccion=request.form.get('direccion', '').strip(),
            comercial=request.form.get('comercial', '').strip(),
            carnet_identidad=request.form.get('carnet_identidad', '').strip(),
            tipo_cliente=request.form.get('tipo_cliente', 'persona'),
            sector=request.form.get('sector', '').strip(),
            preferencias_diseno=request.form.get('preferencias_diseno', '').strip(),
            metodo_pago_favorito=request.form.get('metodo_pago_favorito', '').strip(),
            referido_por=request.form.get('referido_por', '').strip(),
            frecuencia_pedido=request.form.get('frecuencia_pedido', '').strip(),
            notas=request.form.get('notas', '').strip(),
            observaciones_internas=request.form.get('observaciones_internas', '').strip(),
            etiquetas=request.form.get('etiquetas', '').strip(),
            created_by_id=current_user.id
        )
        fn = request.form.get('fecha_nacimiento', '')
        if fn:
            try:
                cliente.fecha_nacimiento = datetime.strptime(fn, '%Y-%m-%d').date()
            except Exception:
                pass

        db.session.add(cliente)
        db.session.commit()
        flash(f'Cliente "{nombre}" creado correctamente.', 'success')
        return redirect(url_for('clientes.editar_cliente', cliente_id=cliente.id) + '#favoritos')

    return render_template('form_cliente.html',
        cliente=None, total_ordenes=0, total_metros=0,
        nivel_actual='standard', color_nivel='#8e8e93', nombre_nivel='Standard',
        etiquetas_favoritas=[], total_favoritos=0)


# ==========================================
# EDITAR CLIENTE
# ==========================================
@clientes_bp.route('/editar/<int:cliente_id>', methods=['GET', 'POST'])
@login_required
@comercial_or_admin_required
def editar_cliente(cliente_id):
    cliente = Client.query.get_or_404(cliente_id)

    if request.method == 'POST':
        cliente.nombre = request.form.get('nombre', '').strip()
        cliente.referencia = request.form.get('referencia', '').strip()
        cliente.telefono = request.form.get('telefono', '').strip()
        cliente.email = request.form.get('email', '').strip()
        cliente.direccion = request.form.get('direccion', '').strip()
        cliente.comercial = request.form.get('comercial', '').strip()
        cliente.carnet_identidad = request.form.get('carnet_identidad', '').strip()
        cliente.tipo_cliente = request.form.get('tipo_cliente', 'persona')
        cliente.sector = request.form.get('sector', '').strip()
        cliente.preferencias_diseno = request.form.get('preferencias_diseno', '').strip()
        cliente.metodo_pago_favorito = request.form.get('metodo_pago_favorito', '').strip()
        cliente.referido_por = request.form.get('referido_por', '').strip()
        cliente.frecuencia_pedido = request.form.get('frecuencia_pedido', '').strip()
        cliente.notas = request.form.get('notas', '').strip()
        cliente.observaciones_internas = request.form.get('observaciones_internas', '').strip()
        cliente.etiquetas = request.form.get('etiquetas', '').strip()

        fn = request.form.get('fecha_nacimiento', '')
        if fn:
            try:
                cliente.fecha_nacimiento = datetime.strptime(fn, '%Y-%m-%d').date()
            except Exception:
                cliente.fecha_nacimiento = None
        else:
            cliente.fecha_nacimiento = None

        if not cliente.nombre:
            flash('El nombre es obligatorio.', 'danger')
            return render_template('form_cliente.html', cliente=cliente, **_get_form_cliente_ctx(cliente))
        if not cliente.referencia:
            cliente.referencia = f"CLI-{datetime.now().strftime('%Y%m%d%H%M%S')}"
        else:
            if Client.query.filter(Client.referencia == cliente.referencia, Client.id != cliente.id).first():
                flash(f'Ya existe otro cliente con la referencia "{cliente.referencia}".', 'danger')
                return render_template('form_cliente.html', cliente=cliente, **_get_form_cliente_ctx(cliente))

        db.session.commit()
        flash(f'Cliente "{cliente.nombre}" actualizado correctamente.', 'success')
        return redirect(url_for('clientes.listar_clientes'))

    return render_template('form_cliente.html', cliente=cliente, **_get_form_cliente_ctx(cliente))


def _get_form_cliente_ctx(cliente):
    """Contexto común para form_cliente (crear/editar)."""
    orders = Order.query.filter_by(client_id=cliente.id).all()
    total_ordenes = len(orders)
    total_metros = 0
    for o in orders:
        for a in o.archivos:
            if a.parametros_etiqueta:
                try:
                    p = json.loads(a.parametros_etiqueta)
                    if p.get('consumo'):
                        total_metros += p.get('consumo', 0)
                        continue
                except Exception:
                    pass
            if a.unidad in ['m', 'm2', 'm²'] and a.cantidad:
                total_metros += a.cantidad
    total_metros = round(total_metros, 2)

    if total_ordenes >= 50 or total_metros >= 500:
        nivel_actual = 'black'
    elif total_ordenes >= 20 or total_metros >= 200:
        nivel_actual = 'golden'
    else:
        nivel_actual = 'standard'

    try:
        config_colores = json.loads(obtener_config('clientes_colores_niveles', '{}'))
    except Exception:
        config_colores = {}
    try:
        config_nombres = json.loads(obtener_config('clientes_nombres_niveles', '{}'))
    except Exception:
        config_nombres = {}

    favoritos = EtiquetaFavorita.query.filter_by(
        client_id=cliente.id, activo=True
    ).order_by(EtiquetaFavorita.nombre.asc()).all()

    return {
        'total_ordenes': total_ordenes,
        'total_metros': total_metros,
        'nivel_actual': nivel_actual,
        'color_nivel': config_colores.get(nivel_actual, '#8e8e93'),
        'nombre_nivel': config_nombres.get(nivel_actual, nivel_actual.capitalize()),
        'etiquetas_favoritas': favoritos,
        'total_favoritos': len(favoritos),
        'favoritos_max_mb': _get_max_favorito_mb(),
        'puede_subir_grande': _puede_subir_archivo_grande(current_user),
    }


# ==========================================
# ELIMINAR CLIENTE
# ==========================================
@clientes_bp.route('/eliminar/<int:cliente_id>', methods=['POST'])
@login_required
@admin_required
def eliminar_cliente(cliente_id):
    cliente = Client.query.get_or_404(cliente_id)
    if cliente.orders and len(cliente.orders) > 0:
        flash(f'No se puede eliminar el cliente "{cliente.nombre}" porque tiene órdenes asociadas.', 'danger')
        return redirect(url_for('clientes.listar_clientes'))
    for fav in cliente.etiquetas_favoritas:
        borrar_archivo_favorito(fav)
    nombre = cliente.nombre
    db.session.delete(cliente)
    db.session.commit()
    flash(f'Cliente "{nombre}" eliminado.', 'success')
    return redirect(url_for('clientes.listar_clientes'))


# ==========================================
# DETALLE DE CLIENTE (SOLO LECTURA DE FAVORITOS)
# ==========================================
@clientes_bp.route('/detalle/<int:cliente_id>')
@login_required
@view_required
def detalle_cliente(cliente_id):
    cliente = Client.query.get_or_404(cliente_id)
    orders = Order.query.filter_by(client_id=cliente_id).all()
    orders_recientes = sorted(orders, key=lambda o: o.date if o.date else datetime.min, reverse=True)[:10]

    total_ordenes = len(orders)
    ordenes_pendientes = sum(1 for o in orders if o.column in ['pendiente', 'por-preparar', 'preparados'])
    ordenes_completadas = sum(1 for o in orders if o.column in ['entregados', 'listo'])

    total_metros = 0
    for o in orders:
        for a in o.archivos:
            if a.parametros_etiqueta:
                try:
                    p = json.loads(a.parametros_etiqueta)
                    if p.get('consumo'):
                        total_metros += p.get('consumo', 0)
                        continue
                except Exception:
                    pass
            if a.unidad in ['m', 'm2', 'm²'] and a.cantidad:
                total_metros += a.cantidad
    total_metros = round(total_metros, 2)

    try:
        config_colores = json.loads(obtener_config('clientes_colores_niveles', '{}'))
    except Exception:
        config_colores = {}
    try:
        config_nombres = json.loads(obtener_config('clientes_nombres_niveles', '{}'))
    except Exception:
        config_nombres = {}

    if total_ordenes >= 50 or total_metros >= 500:
        nivel_actual = 'black'
    elif total_ordenes >= 20 or total_metros >= 200:
        nivel_actual = 'golden'
    else:
        nivel_actual = 'standard'

    insights = calcular_insights_cliente(cliente)

    # ===== Favoritos (solo lectura) =====
    favoritos = EtiquetaFavorita.query.filter_by(
        client_id=cliente_id, activo=True
    ).order_by(
        EtiquetaFavorita.veces_usado.desc(),
        EtiquetaFavorita.ultima_vez_usado.desc().nullslast(),
        EtiquetaFavorita.nombre.asc()
    ).all()

    total_favoritos = len(favoritos)
    total_usos = sum(f.veces_usado or 0 for f in favoritos)
    top_favorito = favoritos[0] if favoritos else None

    # Distribución de uso (para gráfico o lista)
    favoritos_con_uso = [f for f in favoritos if (f.veces_usado or 0) > 0]

    ahora = datetime.now()
    meses, valores = [], []
    for i in range(11, -1, -1):
        mes = ahora.replace(day=1) - timedelta(days=30*i)
        meses.append(mes.strftime('%b %Y'))
        inicio = mes.replace(day=1).date()
        fin = datetime(mes.year + 1, 1, 1).date() if mes.month == 12 else datetime(mes.year, mes.month + 1, 1).date()
        valores.append(sum(1 for o in orders if o.date and inicio <= o.date < fin))

    return render_template('detalle_cliente.html',
        cliente=cliente,
        total_ordenes=total_ordenes,
        total_metros=total_metros,
        ordenes_pendientes=ordenes_pendientes,
        ordenes_completadas=ordenes_completadas,
        frecuencia_pedido=calcular_frecuencia(orders),
        top_servicios=insights['top_servicios'],
        top_materiales=insights['top_materiales'],
        top_productos=insights['top_productos'],
        evolucion_labels=meses, evolucion_values=valores,
        orders_recientes=orders_recientes,
        ultimo_pedido=orders[0].date if orders else None,
        nivel_actual=nivel_actual,
        color_nivel=config_colores.get(nivel_actual, '#8e8e93'),
        nombre_nivel=config_nombres.get(nivel_actual, nivel_actual.capitalize()),
        etiquetas_favoritas=favoritos,
        total_favoritos=total_favoritos,
        total_favoritos_usos=total_usos,
        top_favorito=top_favorito,
        favoritos_con_uso=favoritos_con_uso,
    )


# ==========================================
# ETIQUETAS FAVORITAS · CRUD (desde form_cliente)
# ==========================================
@clientes_bp.route('/<int:cliente_id>/favoritos/crear', methods=['POST'])
@login_required
@comercial_or_admin_required
def crear_favorito(cliente_id):
    cliente = Client.query.get_or_404(cliente_id)

    nombre = (request.form.get('nombre') or '').strip()
    ancho_raw = (request.form.get('ancho_cm') or '').strip()
    alto_raw = (request.form.get('alto_cm') or '').strip()
    notas = (request.form.get('notas') or '').strip()

    archivo = request.files.get('archivo')

    ancho_cm = None
    alto_cm = None
    try:
        if ancho_raw: ancho_cm = float(ancho_raw)
        if alto_raw: alto_cm = float(alto_raw)
    except ValueError:
        pass

    # Si falta algo, intentar auto-extraer del archivo
    if archivo and archivo.filename:
        w_cm, h_cm = _extraer_dimensiones_archivo(archivo)
        if not ancho_cm and w_cm:
            ancho_cm = w_cm
        if not alto_cm and h_cm:
            alto_cm = h_cm
        if not nombre:
            nombre = os.path.splitext(archivo.filename)[0][:150].strip()

    if not nombre:
        flash('Debes indicar un nombre o adjuntar un archivo con nombre.', 'danger')
        return redirect(url_for('clientes.editar_cliente', cliente_id=cliente_id) + '#favoritos')
    if not ancho_cm or not alto_cm or ancho_cm <= 0 or alto_cm <= 0:
        flash('Las dimensiones (ancho y alto en cm) son obligatorias.', 'danger')
        return redirect(url_for('clientes.editar_cliente', cliente_id=cliente_id) + '#favoritos')

    fav = EtiquetaFavorita(
        client_id=cliente_id,
        nombre=nombre,
        ancho_cm=round(ancho_cm, 2),
        alto_cm=round(alto_cm, 2),
        notas=notas,
        created_by_id=current_user.id,
        veces_usado=0,
    )

    if archivo and archivo.filename:
        try:
            ruta, nombre_orig, mime, size = guardar_archivo_favorito(cliente_id, archivo, current_user)
            if ruta:
                fav.archivo_ruta = ruta
                fav.archivo_nombre = nombre_orig
                fav.archivo_tipo = mime
                fav.archivo_tamano = size
        except ValueError as e:
            flash(str(e), 'danger')
            return redirect(url_for('clientes.editar_cliente', cliente_id=cliente_id) + '#favoritos')

    db.session.add(fav)
    db.session.commit()
    flash(f'Etiqueta "{nombre}" añadida.', 'success')
    return redirect(url_for('clientes.editar_cliente', cliente_id=cliente_id) + '#favoritos')


@clientes_bp.route('/favorito/<int:fav_id>/editar', methods=['POST'])
@login_required
@comercial_or_admin_required
def editar_favorito(fav_id):
    fav = EtiquetaFavorita.query.get_or_404(fav_id)

    nombre = (request.form.get('nombre') or '').strip()
    if not nombre:
        flash('El nombre es obligatorio.', 'danger')
        return redirect(url_for('clientes.editar_cliente', cliente_id=fav.client_id) + '#favoritos')

    try:
        ancho_cm = float((request.form.get('ancho_cm') or '').strip())
        alto_cm = float((request.form.get('alto_cm') or '').strip())
    except ValueError:
        flash('Las dimensiones deben ser números.', 'danger')
        return redirect(url_for('clientes.editar_cliente', cliente_id=fav.client_id) + '#favoritos')

    if ancho_cm <= 0 or alto_cm <= 0:
        flash('Las dimensiones deben ser mayores a cero.', 'danger')
        return redirect(url_for('clientes.editar_cliente', cliente_id=fav.client_id) + '#favoritos')

    fav.nombre = nombre
    fav.ancho_cm = round(ancho_cm, 2)
    fav.alto_cm = round(alto_cm, 2)
    fav.notas = (request.form.get('notas') or '').strip()

    # Archivo: puede reemplazarse o eliminarse
    if request.form.get('eliminar_archivo') == 'on':
        borrar_archivo_favorito(fav)
        fav.archivo_ruta = None
        fav.archivo_nombre = None
        fav.archivo_tipo = None
        fav.archivo_tamano = None

    archivo = request.files.get('archivo')
    if archivo and archivo.filename:
        try:
            ruta, nombre_orig, mime, size = guardar_archivo_favorito(fav.client_id, archivo, current_user)
            if ruta:
                borrar_archivo_favorito(fav)
                fav.archivo_ruta = ruta
                fav.archivo_nombre = nombre_orig
                fav.archivo_tipo = mime
                fav.archivo_tamano = size
        except ValueError as e:
            flash(str(e), 'danger')
            return redirect(url_for('clientes.editar_cliente', cliente_id=fav.client_id) + '#favoritos')

    db.session.commit()
    flash(f'Etiqueta "{fav.nombre}" actualizada.', 'success')
    return redirect(url_for('clientes.editar_cliente', cliente_id=fav.client_id) + '#favoritos')


@clientes_bp.route('/favorito/<int:fav_id>/eliminar', methods=['POST'])
@login_required
@comercial_or_admin_required
def eliminar_favorito(fav_id):
    fav = EtiquetaFavorita.query.get_or_404(fav_id)
    client_id = fav.client_id
    nombre = fav.nombre
    borrar_archivo_favorito(fav)
    db.session.delete(fav)
    db.session.commit()
    flash(f'Etiqueta "{nombre}" eliminada.', 'success')
    return redirect(url_for('clientes.editar_cliente', cliente_id=client_id) + '#favoritos')


@clientes_bp.route('/favorito/<int:fav_id>/archivo')
@login_required
@view_required
def ver_archivo_favorito(fav_id):
    fav = EtiquetaFavorita.query.get_or_404(fav_id)
    if not fav.archivo_ruta:
        flash('Esta etiqueta no tiene archivo adjunto.', 'warning')
        return redirect(url_for('clientes.detalle_cliente', cliente_id=fav.client_id) + '#favoritos')
    upload_folder = current_app.config.get('UPLOAD_FOLDER', 'app/static/uploads')
    ruta_absoluta = os.path.join(upload_folder, fav.archivo_ruta)
    if not os.path.exists(ruta_absoluta):
        flash('El archivo no se encuentra en el servidor.', 'danger')
        return redirect(url_for('clientes.detalle_cliente', cliente_id=fav.client_id) + '#favoritos')
    return send_file(
        ruta_absoluta,
        mimetype=fav.archivo_tipo or 'application/octet-stream',
        as_attachment=False,
        download_name=fav.archivo_nombre or 'archivo',
        conditional=True
    )


@clientes_bp.route('/favorito/<int:fav_id>/incrementar-uso', methods=['POST'])
@login_required
@comercial_or_admin_required
def incrementar_uso_favorito(fav_id):
    """Llamado desde el form de órdenes cuando se usa una favorita."""
    fav = EtiquetaFavorita.query.get_or_404(fav_id)
    fav.veces_usado = (fav.veces_usado or 0) + 1
    fav.ultima_vez_usado = datetime.now()
    db.session.commit()
    return jsonify({'success': True, 'veces_usado': fav.veces_usado})


# ==========================================
# EXPORTAR / IMPORTAR / PLANTILLA
# ==========================================
@clientes_bp.route('/exportar')
@login_required
@comercial_or_admin_required
def exportar_clientes():
    clientes = Client.query.order_by(Client.nombre).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Clientes'
    headers = ['ID', 'Referencia', 'Nombre', 'Teléfono', 'Email', 'Dirección',
               'Comercial', 'Carnet Identidad', 'Fecha Nacimiento', 'Tipo Cliente',
               'Sector', 'Preferencias Diseño', 'Método Pago', 'Referido Por',
               'Frecuencia Pedido', 'Notas', 'Observaciones', 'Etiquetas', 'Pedidos']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
    for r, c in enumerate(clientes, 2):
        ws.cell(row=r, column=1, value=c.id)
        ws.cell(row=r, column=2, value=c.referencia)
        ws.cell(row=r, column=3, value=c.nombre)
        ws.cell(row=r, column=4, value=c.telefono or '')
        ws.cell(row=r, column=5, value=c.email or '')
        ws.cell(row=r, column=6, value=c.direccion or '')
        ws.cell(row=r, column=7, value=c.comercial or '')
        ws.cell(row=r, column=8, value=c.carnet_identidad or '')
        ws.cell(row=r, column=9, value=c.fecha_nacimiento.strftime('%Y-%m-%d') if c.fecha_nacimiento else '')
        ws.cell(row=r, column=10, value=c.tipo_cliente or '')
        ws.cell(row=r, column=11, value=c.sector or '')
        ws.cell(row=r, column=12, value=c.preferencias_diseno or '')
        ws.cell(row=r, column=13, value=c.metodo_pago_favorito or '')
        ws.cell(row=r, column=14, value=c.referido_por or '')
        ws.cell(row=r, column=15, value=c.frecuencia_pedido or '')
        ws.cell(row=r, column=16, value=c.notas or '')
        ws.cell(row=r, column=17, value=c.observaciones_internas or '')
        ws.cell(row=r, column=18, value=c.etiquetas or '')
        ws.cell(row=r, column=19, value=len(c.orders) if c.orders else 0)
    for col in range(1, len(headers)+1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True,
                     download_name=f'clientes_{datetime.now().strftime("%Y%m%d")}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@clientes_bp.route('/importar', methods=['GET', 'POST'])
@login_required
@admin_required
def importar_clientes():
    if request.method == 'POST':
        if 'archivo' not in request.files:
            flash('No se seleccionó ningún archivo.', 'danger')
            return redirect(url_for('clientes.importar_clientes'))
        archivo = request.files['archivo']
        if not archivo.filename or not archivo.filename.endswith(('.xlsx', '.xls')):
            flash('Formato no soportado. Use .xlsx o .xls', 'danger')
            return redirect(url_for('clientes.importar_clientes'))
        try:
            wb = openpyxl.load_workbook(archivo)
            ws = wb.active
            headers = [cell.value.strip() if cell.value else '' for cell in ws[1]]
            col_map = {}
            for idx, h in enumerate(headers):
                h_clean = h.lower().replace(' ', '_').replace('(', '').replace(')', '').replace('ñ', 'n')
                for key, needles in {
                    'nombre': ['nombre'], 'referencia': ['referencia'],
                    'telefono': ['telefono'], 'email': ['email'], 'direccion': ['direccion'],
                    'comercial': ['comercial'], 'carnet_identidad': ['carnet', 'identidad'],
                    'fecha_nacimiento': ['fecha_nacimiento', 'fecha nacimiento'],
                    'tipo_cliente': ['tipo_cliente', 'tipo cliente'], 'sector': ['sector'],
                    'preferencias_diseno': ['preferencias'],
                    'metodo_pago_favorito': ['metodo_pago', 'método pago'],
                    'referido_por': ['referido_por', 'referido'],
                    'frecuencia_pedido': ['frecuencia_pedido', 'frecuencia'],
                    'notas': ['notas'], 'observaciones_internas': ['observaciones'],
                    'etiquetas': ['etiquetas'],
                }.items():
                    if any(n in h_clean for n in needles):
                        col_map[key] = idx
                        break
            if 'nombre' not in col_map:
                flash('El archivo debe contener una columna "Nombre".', 'danger')
                return redirect(url_for('clientes.importar_clientes'))
            creados = actualizados = 0
            for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not row or not any(row):
                    continue
                nombre = str(row[col_map['nombre']]).strip() if row[col_map['nombre']] else ''
                if not nombre:
                    continue
                referencia = None
                if col_map.get('referencia') is not None and row[col_map['referencia']]:
                    referencia = str(row[col_map['referencia']]).strip()
                cliente = None
                if referencia:
                    cliente = Client.query.filter_by(referencia=referencia).first()
                if not cliente:
                    cliente = Client.query.filter_by(nombre=nombre).first()
                is_new = cliente is None
                if is_new:
                    cliente = Client(nombre=nombre)
                    cliente.referencia = referencia or f"CLI-{datetime.now().strftime('%Y%m%d%H%M%S')}-{r_idx}"
                    cliente.created_by_id = current_user.id
                for field in ['telefono', 'email', 'direccion', 'comercial', 'carnet_identidad',
                              'tipo_cliente', 'sector', 'preferencias_diseno', 'metodo_pago_favorito',
                              'referido_por', 'frecuencia_pedido', 'notas', 'observaciones_internas', 'etiquetas']:
                    if col_map.get(field) is not None and row[col_map[field]]:
                        setattr(cliente, field, str(row[col_map[field]]).strip())
                if col_map.get('fecha_nacimiento') is not None and row[col_map['fecha_nacimiento']]:
                    try:
                        val = row[col_map['fecha_nacimiento']]
                        if isinstance(val, datetime):
                            cliente.fecha_nacimiento = val.date()
                        elif isinstance(val, (int, float)):
                            cliente.fecha_nacimiento = date(1899, 12, 30) + timedelta(days=int(val))
                        else:
                            cliente.fecha_nacimiento = datetime.strptime(str(val), '%Y-%m-%d').date()
                    except Exception:
                        pass
                if is_new:
                    db.session.add(cliente)
                    creados += 1
                else:
                    actualizados += 1
            db.session.commit()
            flash(f'Importación: {creados} nuevos, {actualizados} actualizados.', 'success')
            return redirect(url_for('clientes.listar_clientes'))
        except Exception as e:
            db.session.rollback()
            flash(f'Error al importar: {str(e)}', 'danger')
            return redirect(url_for('clientes.importar_clientes'))
    return render_template('importar_clientes.html')


@clientes_bp.route('/descargar-plantilla')
@login_required
@admin_required
def descargar_plantilla():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Clientes'
    headers = ['Nombre', 'Referencia', 'Teléfono', 'Email', 'Dirección',
               'Comercial', 'Carnet de Identidad', 'Fecha Nacimiento',
               'Tipo Cliente', 'Sector', 'Preferencias de Diseño',
               'Método Pago Favorito', 'Referido Por', 'Frecuencia Pedido',
               'Notas', 'Observaciones Internas', 'Etiquetas']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(start_color='E8E0D8', end_color='E8E0D8', fill_type='solid')
    for col in range(1, len(headers)+1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 22
    ejemplo = ['Cliente Ejemplo', 'CLI-001', '555-1234', 'cliente@ejemplo.com',
               'Calle Principal 123', 'Comercial 1', '12345678', '1990-01-15',
               'persona', 'Retail', 'Moderno, minimalista', 'Transferencia',
               'Juan Pérez', 'Mensual', 'Cliente frecuente',
               'Prefiere atención personalizada', 'VIP, Premium']
    for col, value in enumerate(ejemplo, 1):
        ws.cell(row=2, column=col, value=value)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='plantilla_clientes.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')